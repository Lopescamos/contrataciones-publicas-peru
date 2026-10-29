"""Exportación a Excel con formato profesional.

* :func:`export_to_excel` / :func:`export_multi_sheet`: cualquier DataFrame a
  Excel con encabezados con estilo, filtros, paneles fijos, anchos de columna y
  formatos numéricos según el nombre de la columna (montos, %, fechas, conteos).
* :func:`report_tables` + :func:`build_report`: reporte ejecutivo del proyecto
  (resumen, rankings, alertas, calidad de datos y notas metodológicas).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from contrataciones.utils.logger import get_logger

logger = get_logger(__name__)

HEADER_LABELS: dict[str, str] = {
    "anio": "Año",
    "categoria": "Categoría",
    "metodo_detalle": "Procedimiento",
    "nombre": "Nombre",
    "ruc": "RUC",
    "departamento": "Departamento",
    "tipo_persona": "Tipo de persona",
    "n_procesos": "N.° procesos",
    "n_procesos_competitivos": "N.° procesos competitivos",
    "n_postor_unico": "N.° con postor único",
    "pct_postor_unico": "% postor único",
    "monto_adjudicado_pen": "Monto adjudicado (S/)",
    "monto_asignado_pen": "Monto adjudicado (S/)",
    "valor_referencial_pen": "Valor referencial (S/)",
    "n_proveedores": "N.° proveedores",
    "n_entidades": "N.° entidades",
    "n_procesos_ganados": "N.° procesos ganados",
    "hhi": "Índice HHI",
    "n_alertas": "N.° alertas",
    "n_procesos_riesgo_alto": "N.° procesos riesgo alto",
    "n_postores": "N.° postores",
    "puntaje_riesgo": "Puntaje de riesgo",
    "nivel_riesgo": "Nivel de riesgo",
    "codigos": "Alertas",
    "detalles": "Detalle de alertas",
    "ocid": "OCID",
    "nomenclatura": "Nomenclatura",
    "entidad": "Entidad",
    "tiene_sancion": "Tiene sanción",
    "codigo": "Código",
    "descripcion": "Descripción",
    "peso": "Peso",
    "referencia": "Referencia",
    "n_procesos_marcados": "N.° procesos marcados",
    "pct_procesos": "% de procesos",
    "indicador": "Indicador",
    "valor": "Valor",
    # Modelo de monto esperado (Fase 2, ML)
    "fecha_convocatoria": "Fecha de convocatoria",
    "monto_esperado_pen": "Monto esperado (S/)",
    "monto_minimo_esperado_pen": "Mínimo esperado (S/)",
    "monto_maximo_esperado_pen": "Máximo esperado (S/)",
    "ratio_adjudicado_referencial": "Adjudicado / referencial",
    "ratio_esperado": "Razón esperada",
    "pct_desviacion": "% desviación",
    "z_robusto": "Z robusto",
    "tipo_desviacion": "Tipo de desviación",
    "es_adjudicado_igual_referencial": "Adjudicado = referencial",
    "origen_prediccion": "Origen de la predicción",
    "n_sobrecosto": "N.° sobrecosto inusual",
    "n_subcosto": "N.° subcosto inusual",
    "n_revisar_datos": "N.° a revisar (datos)",
    "pct_atipicos": "% atípicos",
    "particion": "Partición",
    "modelo": "Modelo",
    "mae_log": "MAE (log)",
    "rmse_log": "RMSE (log)",
    "r2_log_ratio": "R² log(razón)",
    "r2_log_monto": "R² log(monto)",
    "pct_mdape": "MdAPE",
    "pct_wape": "WAPE",
    "variable": "Variable",
    "aumento_mae_log": "Aumento del MAE (log)",
    "desviacion": "Desviación",
    "nota": "Nota",
}


@dataclass
class ExcelStyle:
    header_fill: str = "1F2A44"
    header_font_color: str = "FFFFFF"
    zebra_fill: str = "F4F6F9"
    font_name: str = "Calibri"
    font_size: int = 10
    max_column_width: int = 60

    @classmethod
    def from_config(cls, config: dict[str, Any] | None) -> ExcelStyle:
        config = config or {}
        fields = {k: config[k] for k in cls.__dataclass_fields__ if k in config}
        return cls(**fields)


def infer_number_format(column: str, series: pd.Series | None = None) -> str | None:
    """Formato de Excel según el nombre de la columna."""
    name = column.lower()
    if name.startswith(("pct_", "ratio")):
        return "0.0%"
    if any(token in name for token in ("monto", "valor_referencial", "valor_")) and name != "valor":
        return "#,##0.00"
    if name == "hhi":
        return "#,##0"
    if name == "anio":
        return "0"
    if name.startswith("n_") or name in {
        "puntaje_riesgo",
        "peso",
        "plazo_presentacion_dias",
        "plazo_consultas_dias",
    }:
        return "#,##0"
    if series is not None and pd.api.types.is_datetime64_any_dtype(series):
        return "dd/mm/yyyy"
    return None


def _label(column: str) -> str:
    return HEADER_LABELS.get(column, column.replace("_", " ").capitalize())


def _blank(value: Any) -> bool:
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def apply_formatting(ws: Worksheet, df: pd.DataFrame, header_row: int, style: ExcelStyle) -> None:
    """Estilo de encabezados, filtros, paneles fijos, formatos y anchos."""
    header_font = Font(
        name=style.font_name, size=style.font_size, bold=True, color=style.header_font_color
    )
    header_fill = PatternFill("solid", fgColor=style.header_fill)
    zebra = PatternFill("solid", fgColor=style.zebra_fill)
    body_font = Font(name=style.font_name, size=style.font_size)
    n_rows, n_cols = len(df), len(df.columns)

    for col_idx, column in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        number_format = infer_number_format(str(column), df[column])
        values = ["" if _blank(v) else str(v) for v in df[column].head(500)]
        width = max(
            [len(str(ws.cell(row=header_row, column=col_idx).value or ""))]
            + [len(v) for v in values]
            + [8]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            width + 2, style.max_column_width
        )
        for row in range(header_row + 1, header_row + n_rows + 1):
            body = ws.cell(row=row, column=col_idx)
            body.font = body_font
            if number_format:
                body.number_format = number_format
            if (row - header_row) % 2 == 0:
                body.fill = zebra
    ws.row_dimensions[header_row].height = 30
    if n_cols:
        last = get_column_letter(n_cols)
        ws.auto_filter.ref = f"A{header_row}:{last}{header_row + max(n_rows, 1)}"
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    # Impresión: horizontal y ajustada al ancho de una página.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"


def _write_sheet(
    writer: pd.ExcelWriter,
    name: str,
    df: pd.DataFrame,
    style: ExcelStyle,
    title: str | None = None,
    subtitle: str | None = None,
) -> Worksheet:
    start_row = 3 if title else 0
    renamed = df.rename(columns={c: _label(str(c)) for c in df.columns})
    renamed.to_excel(writer, sheet_name=name[:31], index=False, startrow=start_row)
    ws = writer.sheets[name[:31]]
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(
            name=style.font_name, size=13, bold=True, color=style.header_fill
        )
        if subtitle:
            ws.cell(row=2, column=1, value=subtitle).font = Font(
                name=style.font_name, size=style.font_size, italic=True, color="5B6577"
            )
    # Los formatos se deciden con los nombres técnicos originales de las columnas.
    apply_formatting(ws, df, start_row + 1, style)
    return ws


def export_to_excel(
    df: pd.DataFrame, path: str | Path, sheet_name: str = "Datos", style: ExcelStyle | None = None
) -> Path:
    """Exporta un DataFrame a una hoja de Excel con formato."""
    return export_multi_sheet({sheet_name: df}, path, style=style)


def export_multi_sheet(
    sheets: dict[str, pd.DataFrame],
    path: str | Path,
    style: ExcelStyle | None = None,
    titles: dict[str, tuple[str, str | None]] | None = None,
) -> Path:
    """Exporta varias hojas. ``titles`` = {hoja: (título, subtítulo)}."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    style = style or ExcelStyle()
    titles = titles or {}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            title, subtitle = titles.get(name, (None, None))
            _write_sheet(writer, name, _excel_safe(frame), style, title, subtitle)
    return path


def _excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    """Excel no admite zonas horarias ni tipos nullable con pd.NA en algunas versiones."""
    result = df.copy()
    for column in result.columns:
        series = result[column]
        if isinstance(series.dtype, pd.DatetimeTZDtype):
            result[column] = series.dt.tz_localize(None)
        elif pd.api.types.is_extension_array_dtype(series.dtype) and not isinstance(
            series.dtype, pd.StringDtype
        ):
            result[column] = series.astype(object).where(series.notna(), None)
    return result


def add_bar_chart(
    ws: Worksheet,
    *,
    title: str,
    header_row: int,
    n_rows: int,
    category_col: int,
    value_col: int,
    anchor: str,
    y_title: str = "",
) -> None:
    """Agrega un gráfico de barras nativo de Excel a la hoja."""
    if n_rows == 0:
        return
    chart = BarChart()
    chart.type, chart.style, chart.title = "col", 10, title
    chart.y_axis.title, chart.legend = y_title, None
    chart.y_axis.number_format = '#,##0,," M"'  # millones
    # openpyxl >= 3.1 oculta los ejes por defecto en Excel reciente: se muestran explícitamente.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    data = Reference(ws, min_col=value_col, min_row=header_row, max_row=header_row + n_rows)
    categories = Reference(
        ws, min_col=category_col, min_row=header_row + 1, max_row=header_row + n_rows
    )
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, anchor)


# =============================================================================
# Reporte ejecutivo
# =============================================================================
def report_tables(
    gold: dict[str, pd.DataFrame], top_n: int = 20, max_alerts: int = 1000
) -> dict[str, pd.DataFrame]:
    """Calcula las tablas del reporte a partir de la capa gold."""
    fp = gold["fact_proceso"]
    fa = gold["fact_adjudicacion"]
    entidades = gold["dim_entidad"]
    proveedores = gold["dim_proveedor"]
    metodos = gold["dim_metodo"]
    categorias = gold["dim_categoria"]
    fechas = gold["dim_fecha"][["fecha_key", "anio"]]
    alerts = gold["fact_alerta"]
    indicators = gold["dim_indicador"]

    procesos = (
        fp.merge(fechas, left_on="fecha_convocatoria_key", right_on="fecha_key", how="left")
        .merge(
            metodos[["metodo_key", "metodo_detalle", "es_competitivo"]], on="metodo_key", how="left"
        )
        .merge(categorias[["categoria_key", "categoria"]], on="categoria_key", how="left")
        .merge(
            entidades[["entidad_key", "nombre", "ruc", "departamento"]],
            on="entidad_key",
            how="left",
        )
    )
    procesos["postor_unico"] = procesos["es_competitivo"].fillna(False).astype(bool) & procesos[
        "n_postores"
    ].eq(1).fillna(False).astype(bool)
    procesos["competitivo"] = procesos["es_competitivo"].fillna(False).astype(bool)

    def _summary(by: list[str]) -> pd.DataFrame:
        grouped = (
            procesos.groupby(by, dropna=False)
            .agg(
                n_procesos=("ocid", "count"),
                n_procesos_competitivos=("competitivo", "sum"),
                n_postor_unico=("postor_unico", "sum"),
                monto_adjudicado_pen=("monto_adjudicado_pen", "sum"),
                n_alertas=("n_alertas", "sum"),
            )
            .reset_index()
        )
        denominator = grouped["n_procesos_competitivos"].where(
            grouped["n_procesos_competitivos"] > 0
        )
        grouped["pct_postor_unico"] = (grouped["n_postor_unico"] / denominator).astype("float64")
        return grouped

    por_anio = _summary(["anio"]).sort_values("anio")
    por_categoria = _summary(["categoria"]).sort_values("monto_adjudicado_pen", ascending=False)
    por_metodo = _summary(["metodo_detalle"]).sort_values("n_procesos", ascending=False)

    # Se agrupa por la llave sustituta para no fusionar organizaciones homónimas.
    top_entidades = (
        _summary(["entidad_key", "nombre", "ruc", "departamento"])
        .sort_values("monto_adjudicado_pen", ascending=False)
        .head(top_n)
        .drop(columns="entidad_key")
    )

    valid = fa[fa["es_valida"]].merge(
        proveedores[["proveedor_key", "nombre", "ruc", "tipo_persona", "tiene_sancion"]],
        on="proveedor_key",
        how="inner",
    )
    top_proveedores = (
        valid.groupby(
            ["proveedor_key", "nombre", "ruc", "tipo_persona", "tiene_sancion"], dropna=False
        )
        .agg(
            n_procesos_ganados=("ocid", "nunique"),
            n_entidades=("entidad_key", "nunique"),
            monto_asignado_pen=("monto_asignado_pen", "sum"),
        )
        .reset_index()
        .sort_values("monto_asignado_pen", ascending=False)
        .head(top_n)
        .drop(columns="proveedor_key")
    )

    codes = (
        alerts.groupby("ocid")
        .agg(
            codigos=("codigo", lambda s: ", ".join(sorted(s))),
            detalles=("detalle", lambda s: " | ".join(s)),
        )
        .reset_index()
    )
    alertas = (
        procesos[procesos["n_alertas"] > 0]
        .merge(codes, on="ocid", how="left")
        .rename(columns={"nombre": "entidad"})
        .sort_values(["puntaje_riesgo", "valor_referencial_pen"], ascending=[False, False])
        .head(max_alerts)
    )[
        [
            "ocid",
            "nomenclatura",
            "entidad",
            "departamento",
            "metodo_detalle",
            "anio",
            "valor_referencial_pen",
            "monto_adjudicado_pen",
            "n_postores",
            "puntaje_riesgo",
            "nivel_riesgo",
            "codigos",
            "detalles",
        ]
    ]

    total = max(len(fp), 1)
    flagged = alerts.groupby("codigo")["ocid"].nunique().rename("n_procesos_marcados")
    indicadores = indicators.merge(flagged, left_on="codigo", right_index=True, how="left")
    indicadores["n_procesos_marcados"] = indicadores["n_procesos_marcados"].fillna(0).astype(int)
    indicadores["pct_procesos"] = indicadores["n_procesos_marcados"] / total
    indicadores = indicadores[
        [
            "codigo",
            "nombre",
            "descripcion",
            "peso",
            "n_procesos_marcados",
            "pct_procesos",
            "referencia",
        ]
    ]

    competitive = int(procesos["competitivo"].sum())
    fechas_validas = procesos["fecha_convocatoria_key"].dropna().astype(int)
    periodo = (
        f"{str(fechas_validas.min())[:4]} - {str(fechas_validas.max())[:4]}"
        if not fechas_validas.empty
        else "sin fechas"
    )
    resumen = pd.DataFrame(
        [
            ("Periodo analizado (año de convocatoria)", periodo),
            ("Procesos de contratación", f"{len(fp):,}"),
            ("Entidades compradoras", f"{fp['entidad_key'].nunique():,}"),
            ("Proveedores adjudicados", f"{valid['proveedor_key'].nunique():,}"),
            ("Monto adjudicado total (S/)", f"{procesos['monto_adjudicado_pen'].sum():,.2f}"),
            ("Procesos competitivos", f"{competitive:,}"),
            (
                "% de procesos competitivos con postor único",
                f"{(procesos['postor_unico'].sum() / competitive if competitive else 0):.1%}",
            ),
            ("Procesos con al menos una alerta", f"{int((fp['n_alertas'] > 0).sum()):,}"),
            ("Procesos con nivel de riesgo alto", f"{int((fp['nivel_riesgo'] == 'Alto').sum()):,}"),
            (
                "Proveedores adjudicados con sanción registrada",
                f"{int(valid.loc[valid['tiene_sancion'], 'proveedor_key'].nunique()):,}",
            ),
        ],
        columns=["indicador", "valor"],
    )

    return {
        "Resumen": resumen,
        "Por año": por_anio,
        "Por categoría": por_categoria,
        "Por procedimiento": por_metodo,
        "Top entidades": top_entidades,
        "Top proveedores": top_proveedores,
        "Alertas": alertas,
        "Indicadores": indicadores,
    }


NOTES = [
    "Fuente de datos: OECE (ex OSCE), Portal de Contrataciones Abiertas, estándar OCDS. "
    "Redistribución: Registro de Datos de Open Contracting Partnership (licencia CC BY 4.0).",
    "Los montos se expresan en soles (S/). Los montos en moneda extranjera usan el equivalente en "
    "soles publicado por la fuente o, si no existe, el tipo de cambio del BCRP de la fecha de la "
    "operación.",
    "Los consorcios no tienen RUC propio en la publicación del OECE: se identifican por su código "
    "y se clasifican como 'Consorcio'.",
    "Las alertas de riesgo son señales para priorizar la revisión. NO constituyen prueba de "
    "irregularidad ni de responsabilidad de ninguna entidad o proveedor.",
    "Metodología de indicadores: docs/08_indicadores_de_riesgo.md (basada en la guía 'Red Flags "
    "in Public Procurement', OCP 2024, y en la librería Cardinal).",
    "La publicación OCDS no incluye las contrataciones sin procedimiento de selección, como las "
    "menores a 3, 8 o 9 UIT según la normativa aplicable.",
]


def build_report(
    gold: dict[str, pd.DataFrame],
    quality: pd.DataFrame | None,
    path: str | Path,
    *,
    top_n: int = 20,
    max_alerts: int = 1000,
    style: ExcelStyle | None = None,
    source_label: str = "",
) -> Path:
    """Genera el reporte ejecutivo en Excel."""
    style = style or ExcelStyle()
    tables = report_tables(gold, top_n=top_n, max_alerts=max_alerts)
    if quality is not None:
        tables["Calidad de datos"] = quality
    tables["Notas"] = pd.DataFrame({"nota": NOTES})
    generated = datetime.now().strftime("%d/%m/%Y %H:%M")
    subtitle = f"Generado el {generated}" + (f" | Fuente: {source_label}" if source_label else "")
    titles = {
        "Resumen": ("Contrataciones públicas del Perú: resumen ejecutivo", subtitle),
        "Por año": ("Evolución anual", "Monto adjudicado y competencia por año de convocatoria"),
        "Por categoría": ("Gasto por categoría", "Bienes, servicios, obras y consultoría"),
        "Por procedimiento": (
            "Procedimientos de selección",
            "Competencia y monto por procedimiento",
        ),
        "Top entidades": (f"Top {top_n} entidades por monto adjudicado", None),
        "Top proveedores": (f"Top {top_n} proveedores por monto adjudicado", None),
        "Alertas": (
            "Procesos con alertas de riesgo",
            "Ordenados por puntaje. Una alerta no prueba una irregularidad.",
        ),
        "Indicadores": ("Catálogo de indicadores de riesgo", None),
        "Calidad de datos": (
            "Reporte de calidad de datos",
            "Reglas evaluadas sobre la capa silver",
        ),
        "Notas": ("Notas metodológicas", None),
    }
    path = export_multi_sheet(tables, path, style=style, titles=titles)

    # Gráfico de barras del monto por año (hoja "Por año").
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    ws = workbook["Por año"]
    columns = list(tables["Por año"].columns)
    add_bar_chart(
        ws,
        title="Monto adjudicado por año (S/)",
        header_row=4,
        n_rows=len(tables["Por año"]),
        category_col=columns.index("anio") + 1,
        value_col=columns.index("monto_adjudicado_pen") + 1,
        anchor=f"{get_column_letter(len(columns) + 2)}4",
    )
    workbook.save(path)
    logger.info("Reporte Excel generado: %s", path)
    return path
