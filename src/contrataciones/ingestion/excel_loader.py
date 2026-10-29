"""Lectura de Excel y CSV con detección de encabezados y mapeo de columnas.

Las fuentes públicas cambian sus encabezados con frecuencia ("RUC",
"Nro. RUC", "RUC del proveedor"...). Este módulo normaliza los nombres y los
mapea a columnas canónicas usando los alias de ``config/sources.yaml``.

Ejemplo:
    >>> mapping = {"ruc": ["NRO_RUC", "RUC"], "nombre": ["RAZON_SOCIAL"]}
    >>> df = load_tabular("sanciones.xlsx", header_row="auto", aliases=mapping)
    >>> df = apply_column_mapping(df, mapping, required=["ruc"])
"""

from __future__ import annotations

import csv
import re
import unicodedata
from collections.abc import Iterable, Mapping
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from contrataciones.utils.logger import get_logger

logger = get_logger(__name__)

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
CSV_SUFFIXES = {".csv", ".txt", ".tsv"}


class SchemaError(ValueError):
    """El archivo no tiene las columnas requeridas."""


def normalize_header(name: Any) -> str:
    """Normaliza un encabezado: MAYÚSCULAS, sin tildes, separadores -> ``_``.

    >>> normalize_header("  Número de Resolución ")
    'NUMERO_DE_RESOLUCION'
    >>> normalize_header("Nº RUC")
    'NO_RUC'
    """
    text = unicodedata.normalize("NFKD", str(name))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^0-9A-Za-z]+", "_", text.upper())
    return text.strip("_")


def list_sheets(path: str | Path) -> list[str]:
    """Nombres de las hojas de un libro Excel."""
    workbook = load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _alias_keys(aliases: Mapping[str, Iterable[str]] | Iterable[str] | None) -> set[str]:
    if not aliases:
        return set()
    if isinstance(aliases, Mapping):
        values: list[str] = []
        for canonical, names in aliases.items():
            values.append(canonical)
            values.extend(names)
        return {normalize_header(v) for v in values}
    return {normalize_header(v) for v in aliases}


def detect_header_row(raw: pd.DataFrame, aliases: Any, max_rows: int = 20) -> int:
    """Devuelve el índice de la fila (0-based) con más coincidencias de alias.

    Útil cuando el archivo trae títulos o notas encima de la tabla.
    """
    keys = _alias_keys(aliases)
    best_row, best_hits = 0, 0
    for index in range(min(max_rows, len(raw))):
        cells = {normalize_header(v) for v in raw.iloc[index].tolist() if pd.notna(v)}
        hits = len(cells & keys)
        if hits > best_hits:
            best_row, best_hits = index, hits
    return best_row


def _finalize(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.dropna(how="all")
    frame = frame.loc[
        :, [str(c).strip() != "" and not str(c).startswith("Unnamed") for c in frame.columns]
    ]
    frame.columns = [str(c).strip() for c in frame.columns]
    return frame.reset_index(drop=True)


def load_excel(
    path: str | Path,
    sheet: int | str = 0,
    header_row: int | str = 0,
    aliases: Any = None,
) -> pd.DataFrame:
    """Lee una hoja de Excel como texto. ``header_row='auto'`` detecta la fila de encabezados."""
    path = Path(path)
    if header_row == "auto":
        raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=str, nrows=25)
        header_row = detect_header_row(raw, aliases)
        logger.debug("Encabezados detectados en la fila %d de %s", header_row, path.name)
    frame = pd.read_excel(path, sheet_name=sheet, header=int(header_row), dtype=str)
    return _finalize(frame)


def _sniff_separator(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;|\t").delimiter
    except csv.Error:
        return ","


def load_csv(
    path: str | Path,
    sep: str | None = None,
    encoding: str | None = None,
    header_row: int | str = 0,
    aliases: Any = None,
) -> pd.DataFrame:
    """Lee un CSV como texto, detectando separador y codificación (UTF-8 o Latin-1)."""
    path = Path(path)
    encodings = [encoding] if encoding else ["utf-8-sig", "latin-1"]
    last_error: Exception | None = None
    for candidate in encodings:
        try:
            with path.open(encoding=candidate) as handle:
                sample = handle.read(20000)
            separator = sep or _sniff_separator(sample)
            if header_row == "auto":
                raw = pd.read_csv(
                    path, sep=separator, encoding=candidate, header=None, dtype=str, nrows=25
                )
                header = detect_header_row(raw, aliases)
            else:
                header = int(header_row)
            frame = pd.read_csv(path, sep=separator, encoding=candidate, header=header, dtype=str)
            return _finalize(frame)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"No se pudo leer {path} con {encodings}: {last_error}")


def load_tabular(path: str | Path, **kwargs: Any) -> pd.DataFrame:
    """Lee Excel o CSV según la extensión del archivo."""
    suffix = Path(path).suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        return load_excel(
            path, **{k: v for k, v in kwargs.items() if k in {"sheet", "header_row", "aliases"}}
        )
    if suffix in CSV_SUFFIXES:
        return load_csv(
            path,
            **{
                k: v for k, v in kwargs.items() if k in {"sep", "encoding", "header_row", "aliases"}
            },
        )
    raise ValueError(f"Formato no soportado para {path} (use .xlsx, .xls o .csv)")


def validate_columns(df: pd.DataFrame, expected: Iterable[str]) -> list[str]:
    """Devuelve las columnas esperadas que faltan en el DataFrame."""
    present = set(df.columns)
    return [column for column in expected if column not in present]


def apply_column_mapping(
    df: pd.DataFrame,
    mapping: Mapping[str, Iterable[str]],
    required: Iterable[str] = (),
    keep_unmapped: bool = False,
) -> pd.DataFrame:
    """Renombra columnas a su nombre canónico usando alias normalizados.

    Raises:
        SchemaError: si falta alguna columna requerida.
    """
    normalized = {normalize_header(column): column for column in df.columns}
    rename: dict[str, str] = {}
    for canonical, aliases in mapping.items():
        for alias in (canonical, *aliases):
            original = normalized.get(normalize_header(alias))
            if original is not None and original not in rename:
                rename[original] = canonical
                break
    missing = [column for column in required if column not in rename.values()]
    if missing:
        raise SchemaError(
            f"Faltan columnas requeridas {missing}. Encabezados encontrados: {list(df.columns)}. "
            "Ajusta los alias en config/sources.yaml."
        )
    renamed = df.rename(columns=rename)
    if keep_unmapped:
        return renamed
    columns = [canonical for canonical in mapping if canonical in renamed.columns]
    result = renamed[columns].copy()
    for canonical in mapping:
        if canonical not in result.columns:
            result[canonical] = pd.NA
    return result[list(mapping)]
