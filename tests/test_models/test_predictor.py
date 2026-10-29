"""Pruebas del puntaje de los procesos, la exportación y el reporte Excel."""

from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from contrataciones.ingestion.db_connector import DatabaseConnector
from contrataciones.models import anomaly, predictor, report

TODAY = date(2026, 1, 15)


@pytest.fixture(scope="module")
def predictions(trained):
    return predictor.score_processes(trained.artifact, trained.gold, today=TODAY)


def test_output_schema_and_population(predictions, trained):
    assert list(predictions.columns) == list(predictor.OUTPUT_COLUMNS)
    assert len(predictions) == int(trained.table["es_poblacion"].sum())
    assert predictions["ocid"].is_unique
    assert set(predictions["tipo_desviacion"]) <= set(anomaly.DEVIATION_LABELS)
    assert (predictions["fecha_prediccion"] == pd.Timestamp(TODAY)).all()
    assert (predictions["modelo_version"] == trained.artifact["version"]).all()


def test_history_uses_out_of_fold_predictions(predictions, trained):
    assert (predictions["origen_prediccion"] == predictor.ORIGIN_HISTORICAL).all()
    stored = trained.artifact["oof"]
    expected = np.exp(predictions["ocid"].map(stored).astype("float64"))
    assert (predictions["ratio_esperado"] - expected).abs().max() < 1e-6


def test_planted_anomalies_and_data_errors_are_flagged(predictions):
    counts = predictions["tipo_desviacion"].value_counts()
    assert counts[anomaly.LABEL_DATA_ERROR] == 5
    assert counts[anomaly.LABEL_OVERCOST] >= 6  # 8 sobrecostos plantados de x2.2
    assert counts[anomaly.LABEL_UNDERCOST] >= 6  # 8 subcostos plantados de x0.45
    flagged = predictions[predictions["es_atipico"]]
    over = flagged[flagged["tipo_desviacion"] == anomaly.LABEL_OVERCOST]
    assert (over["monto_adjudicado_pen"] > over["monto_maximo_esperado_pen"]).all()
    assert (over["pct_desviacion"] >= 0.15).all()


def test_new_processes_use_the_final_model(trained):
    artifact = {**trained.artifact, "oof": trained.artifact["oof"].iloc[:-50]}
    scored = predictor.score_processes(artifact, trained.gold, today=TODAY)
    origins = scored["origen_prediccion"].value_counts()
    assert origins[predictor.ORIGIN_NEW] == 50
    summary = predictor.summarize(scored)
    assert summary["por_origen"][predictor.ORIGIN_NEW] == 50
    assert sum(summary["por_tipo"].values()) == summary["procesos"]


def test_export_predictions(tmp_path, predictions):
    written = predictor.export_predictions(predictions, tmp_path, "pred_monto_esperado")
    assert set(written) == {"parquet", "csv"}
    back = pd.read_parquet(written["parquet"])
    assert len(back) == len(predictions)
    assert predictor.export_predictions(predictions, tmp_path, "x", formats=["csv"]).keys() == {
        "csv"
    }


def test_load_to_database_without_gold_skips_views(tmp_path, predictions, caplog):
    config = {
        "active_profile": "prueba",
        "profiles": {"prueba": {"dialect": "sqlite", "database": str(tmp_path / "ml.db")}},
    }
    sql_dir = Path(__file__).resolve().parents[2] / "sql"
    result = predictor.load_to_database(
        predictions, table="pred_monto_esperado", sql_dir=sql_dir, db_config=config
    )
    assert result["filas"] == len(predictions)
    assert result["vistas"] == []
    assert "no se crean las vistas" in caplog.text
    # Una segunda carga reemplaza (full refresh), no duplica.
    predictor.load_to_database(
        predictions, table="pred_monto_esperado", sql_dir=sql_dir, db_config=config
    )
    with DatabaseConnector(config=config) as db:
        count = db.read_query("SELECT COUNT(*) AS n FROM pred_monto_esperado")["n"].item()
    assert count == len(predictions)


def test_excel_report(tmp_path, predictions, trained):
    path = report.write_report(
        predictions, trained.artifact, trained.gold, tmp_path / "reporte.xlsx", max_rows=10
    )
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Resumen",
        "Atípicos",
        "Por categoría",
        "Métricas",
        "Importancia",
        "Notas",
    ]
    tables = report.report_tables(predictions, trained.artifact, trained.gold, max_rows=10)
    outliers = tables["Atípicos"]
    assert 0 < len(outliers) <= 10
    assert outliers["tipo_desviacion"].iloc[0] == anomaly.LABEL_OVERCOST
    assert {"entidad", "metodo_detalle", "categoria", "fecha_convocatoria"} <= set(outliers.columns)
    summary = tables["Resumen"].set_index("indicador")["valor"]
    assert summary["Modelo aceptado"] == "Sí"
    assert summary["Procesos evaluados"] == len(predictions)
